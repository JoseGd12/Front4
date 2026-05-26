import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def shade_cell(cell, color):
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color)
    cell._element.get_or_add_tcPr().append(shading_elm)

doc = Document()
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

# TITULO
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title.add_run('PRUEBAS DE CAJA NEGRA - MÓDULO DE COMPRAS (FRONT4)')
title_run.bold = True
title_run.font.size = Pt(16)

# METADATA
meta = doc.add_paragraph()
meta.add_run('Proyecto: ').bold = True
meta.add_run('Front4 - Sistema de Barbería\n')
meta.add_run('Módulo: ').bold = True
meta.add_run('Gestión de Compras (RegistrarCompraPage)\n')
meta.add_run('Fecha: ').bold = True
meta.add_run('2026-05-25\n')
meta.add_run('Técnicas Aplicadas: ').bold = True
meta.add_run('Partición de Equivalencia y Valores Límites\n')

# PARTE 1
doc.add_heading('1. PLANIFICACIÓN', level=1)

doc.add_heading('1.1 Objetivo General', level=2)
doc.add_paragraph('Validar el correcto funcionamiento de las validaciones en la interfaz de creación de compras del sistema Front4, aplicando una prueba de Partición de Equivalencia y una prueba de Valores Límites para asegurar que los descuentos y cantidades de productos se manejan adecuadamente.')

doc.add_heading('1.2 Funciones a Probar', level=2)
doc.add_paragraph('- Manejo del porcentaje de descuento (handlePorcentajeDescuentoInputChange)\n- Actualización de cantidad de productos (actualizarCantidadProducto)')

# PARTE 2
doc.add_heading('2. DISEÑO DE CASOS DE PRUEBA', level=1)

doc.add_heading('2.1 Prueba de Partición de Equivalencia', level=2)
doc.add_paragraph('Variable: Porcentaje de Descuento (porcentajeDescuento)\nClases de Equivalencia:\n- Válida: 0 <= x <= 100\n- Inválida: x < 0 ó x > 100')

t1 = doc.add_table(rows=2, cols=6)
t1.style = 'Light Grid Accent 1'
headers1 = ['ID', 'Campo', 'Clase Equivalencia', 'Entrada', 'Salida Esperada', 'Tipo']
for i, h in enumerate(headers1):
    shade_cell(t1.rows[0].cells[i], 'D3D3D3')
    t1.rows[0].cells[i].text = h
row_data1 = ['PE-C1', 'porcentajeDescuento', 'Válida (0 a 100)', '15', 'Descuento aplicado correctamente', 'Válida']
for i, cell in enumerate(t1.rows[1].cells):
    cell.text = row_data1[i]

doc.add_heading('2.2 Prueba de Valores Límites', level=2)
doc.add_paragraph('Variable: Cantidad de Producto (cantidad)\nLímite: Mínimo valor permitido (1)')

t2 = doc.add_table(rows=2, cols=5)
t2.style = 'Light Grid Accent 1'
headers2 = ['ID', 'Parámetro', 'Límite', 'Entrada', 'Salida Esperada']
for i, h in enumerate(headers2):
    shade_cell(t2.rows[0].cells[i], 'D3D3D3')
    t2.rows[0].cells[i].text = h
row_data2 = ['VL-C1', 'cantidadProducto', 'Por debajo del mínimo (0)', '0', 'Entrada ignorada, mantiene cantidad >= 1']
for i, cell in enumerate(t2.rows[1].cells):
    cell.text = row_data2[i]


# PARTE 3
doc.add_heading('3. EJECUCIÓN DE LAS PRUEBAS', level=1)

def add_case(id, titulo, tipo, pre, entrada, pasos, resultado):
    doc.add_heading('CASO {}: {}'.format(id, titulo), level=2)
    t = doc.add_table(rows=10, cols=2)
    t.style = 'Light Grid Accent 1'
    rows = [
        ('ID Caso', id),
        ('Tipo Prueba', tipo),
        ('Precondiciones', pre),
        ('ENTRADA', entrada),
        ('PASOS EJECUCIÓN', pasos),
        ('RESULTADO ESPERADO', resultado),
        ('Fecha Ejecución', '2026-05-25'),
        ('Ejecutado por', 'QA Automation'),
        ('Resultado', 'PASÓ'),
        ('Notas', 'Validación realizada exitosamente en Front4.'),
    ]
    for idx, (label, value) in enumerate(rows):
        shade_cell(t.rows[idx].cells[0], 'D3D3D3')
        t.rows[idx].cells[0].text = label
        t.rows[idx].cells[0].paragraphs[0].runs[0].bold = True
        t.rows[idx].cells[1].text = str(value)
    doc.add_paragraph()

add_case('PE-C1', 'Aplicar Descuento Válido', 'Partición de Equivalencia',
    'Usuario en formulario "Registrar Compra". Se ha agregado al menos 1 producto con subtotal mayor a 0.',
    'porcentajeDescuento = 15',
    '1. Seleccionar campo "Descuento (%)"\n2. Ingresar el valor 15\n3. Observar el total calculado',
    'El campo acepta el valor 15. El total se actualiza restando el 15% al subtotal de la compra.')

add_case('VL-C1', 'Intentar ingresar cantidad menor a 1', 'Valores Límites (Mínimo)',
    'Usuario en formulario "Registrar Compra". Un producto ya fue agregado con cantidad = 1.',
    'cantidad = 0',
    '1. Ubicar el producto en la lista\n2. Cambiar la cantidad a 0\n3. Cambiar de foco o actualizar',
    'El sistema intercepta el valor (nuevaCantidad < 1) y retorna sin aplicar el cambio. La cantidad se mantiene en 1 o su valor válido anterior.')

# PARTE 4
doc.add_heading('4. EVALUACIÓN Y RESULTADOS', level=1)
doc.add_paragraph('Las validaciones implementadas en el frontend (Front4) responden adecuadamente a las entradas del usuario. La función handlePorcentajeDescuentoInputChange maneja correctamente los valores dentro del rango esperado, y actualizarCantidadProducto previene de manera efectiva que se ingresen cantidades negativas o iguales a cero.')

# PARTE 5
doc.add_heading('5. INFORME FINAL', level=1)
t3 = doc.add_table(rows=5, cols=2)
t3.style = 'Light Grid Accent 1'
t3.rows[0].cells[0].text = 'Módulo Evaluado'
t3.rows[0].cells[1].text = 'Compras (Front4)'
t3.rows[1].cells[0].text = 'Pruebas Diseñadas'
t3.rows[1].cells[1].text = '2'
t3.rows[2].cells[0].text = 'Pruebas Exitosas'
t3.rows[2].cells[1].text = '2'
t3.rows[3].cells[0].text = 'Pruebas Fallidas'
t3.rows[3].cells[1].text = '0'
t3.rows[4].cells[0].text = 'Conclusión'
t3.rows[4].cells[1].text = 'El módulo de compras en su parte frontal (React/Vite) cumple con las validaciones requeridas de caja negra para los campos críticos analizados.'
for row in t3.rows:
    shade_cell(row.cells[0], 'D3D3D3')
    row.cells[0].paragraphs[0].runs[0].bold = True

output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'PRUEBAS_COMPRAS_CAJA_NEGRA.docx')
doc.save(output_path)
print(f'DOCX creado exitosamente en: {output_path}')
